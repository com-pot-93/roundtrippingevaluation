import os

import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook

CHECK_DIR = './check'
PICS_DIR = './pics'
PREFIX = 'text_model_'
SUFFIX = '.xlsx'

PRECISION_COLOR = '#1a5fb4'
RECALL_COLOR = '#26a269'
BAR_HEIGHT = 0.7

DATASET_ORDER = ['domain', 'mad', 'pet', 'sapsam', 'realset']
MISSING = (float('nan'), float('nan'))


def average(values):
    numeric = [v for v in values if isinstance(v, (int, float))]
    return sum(numeric) / len(numeric) if numeric else float('nan')


def load_dataset_averages(path):
    """(average precision, average recall) over the file rows of one text_model_<dataset>.xlsx,
    with the columns looked up by header name."""
    workbook = load_workbook(path, read_only=True)
    header, *rows = workbook.active.iter_rows(values_only=True)
    workbook.close()

    rows = [row for row in rows if row[0]]
    precision_col, recall_col = header.index('precision'), header.index('recall')
    return (
        average([row[precision_col] for row in rows]),
        average([row[recall_col] for row in rows]),
    )


def load_folder_averages(folder):
    averages = {}
    for filename in sorted(os.listdir(folder)):
        if filename.startswith(PREFIX) and filename.endswith(SUFFIX):
            averages[filename[len(PREFIX):-len(SUFFIX)]] = load_dataset_averages(os.path.join(folder, filename))
    return averages


def order_datasets(datasets):
    ordered = [d for d in DATASET_ORDER if d in datasets]
    return ordered + sorted(set(datasets) - set(ordered))


def find_thresholds(check_dir):
    thresholds = []
    for name in os.listdir(check_dir):
        if not os.path.isdir(os.path.join(check_dir, name)):
            continue
        try:
            thresholds.append((float(name), name))
        except ValueError:
            continue
    return [name for _, name in sorted(thresholds)]


def label(value):
    return '' if np.isnan(value) else '{:.2f}'.format(value)


def plot_text_model_bars(check_dir=CHECK_DIR):
    """One figure, one subplot per threshold folder. Per dataset a single row: the recall bar drawn
    over the precision bar, both the same thickness (recall is the smaller one, so it stays visible)."""
    thresholds = find_thresholds(check_dir)
    averages = {t: load_folder_averages(os.path.join(check_dir, t)) for t in thresholds}
    datasets = order_datasets({d for folder in averages.values() for d in folder})
    if not datasets:
        print('No {}<dataset>{} files found in the threshold subfolders of {}.'.format(PREFIX, SUFFIX, check_dir))
        return

    y = np.arange(len(datasets))
    fig, axes = plt.subplots(1, len(thresholds), figsize=(6 * len(thresholds), 5),
                             sharex=True, sharey=True, squeeze=False)
    axes = axes[0]

    for ax, threshold in zip(axes, thresholds):
        precision = [averages[threshold].get(d, MISSING)[0] for d in datasets]
        recall = [averages[threshold].get(d, MISSING)[1] for d in datasets]

        precision_bars = ax.barh(y, np.nan_to_num(precision), BAR_HEIGHT, color=PRECISION_COLOR, label='Precision')
        recall_bars = ax.barh(y, np.nan_to_num(recall), BAR_HEIGHT, color=RECALL_COLOR, label='Recall')
        ax.bar_label(precision_bars, labels=[label(v) for v in precision], padding=3, fontsize=8)
        ax.bar_label(recall_bars, labels=[label(v) for v in recall], label_type='center', color='white', fontsize=8)

        ax.set_xlim(0, 1.1)
        ax.set_xlabel('Average score')
        ax.set_title('Threshold {}'.format(threshold))

    axes[0].set_yticks(y)
    axes[0].set_yticklabels(datasets)
    axes[0].set_ylabel('Dataset')
    axes[0].invert_yaxis()

    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc='upper center', bbox_to_anchor=(0.5, 0.0), ncol=2, frameon=False)
    plt.tight_layout()

    os.makedirs(PICS_DIR, exist_ok=True)
    base = os.path.join(PICS_DIR, 'text_model_bars')
    plt.savefig(base + '.png', dpi=300, bbox_inches='tight')
    plt.savefig(base + '.pdf', bbox_inches='tight')
    plt.savefig(base + '.svg', bbox_inches='tight')
    plt.close(fig)
    print('Saved {}.png, {}.pdf and {}.svg'.format(base, base, base))


if __name__ == "__main__":
    if not os.path.isdir(CHECK_DIR):
        print('Folder {} not found. Run text_model_evaluation.py first.'.format(CHECK_DIR))
    elif not find_thresholds(CHECK_DIR):
        print('No threshold subfolders (e.g. 0.6) found in {}.'.format(CHECK_DIR))
    else:
        plot_text_model_bars()
